"""Tier C. GeM — the primary channel, and our strongest AI story.

There is no public seller API. The category-specific Excel bulk upload IS the integration,
and it is an official documented path rather than a workaround: fill the category's Excel
accurately, save, and the item is listed on GeM 3.0.

🔑 Why this is the part that is genuinely hard, and therefore genuinely valuable:

GeM has over 10,700 product categories. Choosing the wrong one is the single most common
listing failure — an entire consultancy industry exists purely to fill GeM catalogues
correctly. An artisan cannot navigate 10,700 categories. Neither can most educated sellers.

Our vision model maps a photo of a Sambalpuri saree to the right GeM category, the NLP
fills the attribute schema, and we look up the HSN code. That is a consultancy industry,
automated, free, in the artisan's own language. Background removal is a commodity; this
is not.

⚠️ BLOCKING: the real category templates are spec §18 items 1 and 2 and are not in the
repo yet. `_load_template` reads a JSON descriptor per category so that when the real
templates arrive, adding one is dropping in a file — never editing this adapter.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

from openpyxl import Workbook

from .base import ChannelAdapter, PublishResult, Tier

TEMPLATE_DIR = Path(__file__).parent / "gem_templates"


def _load_template(gem_category_id: str | None) -> dict | None:
    if not gem_category_id:
        return None
    path = TEMPLATE_DIR / f"{gem_category_id}.json"
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


class GeMAdapter(ChannelAdapter):
    id = "gem"
    name = "GeM"
    tier = Tier.C

    def map_category(self, product) -> dict:
        return {"gem_id": (product.category_map or {}).get("gem_id")}

    def map_attributes(self, product) -> dict:
        """Fill the GeM field set.

        `local_content` and `country_of_origin` are required for an OEM, and an artisan
        who makes their own goods IS an OEM rather than a Reseller — getting that wrong is
        a rejection on its own.
        """
        artisan = product.artisan
        return {
            "product_name": product.title,
            "description": product.desc_en,
            "brand": artisan.display_name,
            "brand_type": "Unregistered",
            "seller_type": "OEM",
            "material": product.material,
            "hsn_code": product.hsn_code,
            "sku_id": product.id,
            "country_of_origin": "India",
            "local_content_percent": 100,
            "mrp": float(product.mrp or 0),
            "offer_price": float(product.price or 0),
            "certifications": ", ".join(product.certifications or []),
            "gi_claim": product.gi_claim or "",
            "dimensions": product.dimensions,
            "weight_grams": product.weight_grams,
        }

    def format_images(self, product) -> list[dict]:
        # GeM wants 3 images, plain white background, no watermark, no packaging.
        # ⚠️ Exact pixel specs are spec §18 item 4 — unconfirmed. Do not put a number on
        # a slide until it comes from GeM directly.
        usable = [i for i in product.images if not i.is_generated]
        return [{"url": i.url} for i in usable[:3]]

    def check_discount(self, product) -> str | None:
        """🚨 GeM mandates a minimum discount off MRP when listing (~10% typical).

        The floor-price guard has to run AFTER that discount, not before. Suggest ₹2,000,
        GeM knocks 10% off, the artisan nets ₹1,800 — which may be below what it cost to
        make. Miss this and we recommend loss-making prices on the platform we headline.

        Priced upstream in ai/price (the MRP is set so the post-discount price still
        clears the floor); this is the assertion that it actually happened.
        """
        if not product.mrp or not product.floor_price:
            return "missing mrp or floor price"
        post_discount = float(product.mrp) * 0.9
        if post_discount < float(product.floor_price):
            return (
                f"after GeM's 10% mandated discount the price is {post_discount:.0f}, "
                f"below the floor of {float(product.floor_price):.0f}"
            )
        return None

    def build_workbook(self, product) -> tuple[bytes, list[str]]:
        """Render the category Excel. Returns (bytes, warnings)."""
        warnings: list[str] = []
        gem_id = self.map_category(product)["gem_id"]
        template = _load_template(gem_id)
        attrs = self.map_attributes(product)

        if template is None:
            warnings.append(
                f"no GeM template for category {gem_id!r} — emitting a generic sheet. "
                "See docs/Application-Architecture.md §12: real templates are blocking."
            )
            columns = [{"header": k, "field": k} for k in attrs]
        else:
            columns = template["columns"]

        wb = Workbook()
        ws = wb.active
        ws.title = (template or {}).get("sheet_name", "Catalog")
        ws.append([c["header"] for c in columns])

        row = []
        for col in columns:
            value = attrs.get(col["field"])
            if value in (None, "") and col.get("required"):
                warnings.append(f"required GeM field is empty: {col['header']}")
            row.append(value)
        ws.append(row)

        for i, img in enumerate(self.format_images(product), start=1):
            ws.cell(row=1, column=len(columns) + i, value=f"Image{i}")
            ws.cell(row=2, column=len(columns) + i, value=img["url"])

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), warnings

    async def render(self, product, status) -> PublishResult:
        if bad := self.preflight(product):
            return bad

        if problem := self.check_discount(product):
            # Refuse rather than generate a file that loses the artisan money. This is the
            # single feature that does the most for the "increase annual income" goal —
            # under-pricing is the epidemic in this sector, not over-pricing.
            return PublishResult(
                channel=self.id,
                status="failed",
                message_key="price.floor_warning",
                error=problem,
            )

        data, warnings = self.build_workbook(product)
        # TODO(phase 7): persist to object storage and hand back a signed URL. The artisan
        # (or their cluster coordinator) uploads it — GeM has no API to push it for them,
        # and pretending otherwise is a question we cannot survive.
        return PublishResult(
            channel=self.id,
            status="file_ready",
            artifact_url=f"s3://{self.id}/{product.id}.xlsx",
            message_key="publish.file_ready",
            error="; ".join(warnings) or None,
            instructions=[{"voice_key": "gem.step.upload", "copy": product.id}],
        )
